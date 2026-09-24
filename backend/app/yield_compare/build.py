"""전국 2021–2025 수익률 비교 스냅샷.

backend에서:
  python -m app.yield_compare.build
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text

from app.built.db import get_built_engine
from app.collective.db import get_collective_engine
from app.macro_ts.annual_scale import _latest
from app.regional_profile.ecos_csv import find_named_series, parse_ecos_calendar_years
from app.rent.db import get_rent_engine
from app.rent.sangkwon_agg import compound_annual, parse_number, parse_quarter_header
from app.yield_compare.price_pdf import december_index, december_return_pct

ROOT = Path(__file__).resolve().parents[3]
DATA = ROOT / "data"
ANNUAL_DIR = DATA / "한은 연간"
SANGKWON_DIR = ROOT / "임대시장" / "B.상업용"
OUT = ROOT / "docs" / "lab" / "yield_compare_national.json"

YEARS = (2021, 2022, 2023, 2024, 2025)
# 전국 연간 가운데값. 이보다 적으면 그 해는 빈칸.
MIN_N = 100
EXPENSE_RATIO = 0.10

PRICE_PDF = {
    "apartment": "아파트_전국 매매가격지수.pdf",
    "rowhouse": "연립다세대_전국 매매가격지수.pdf",
    "officetel": "오피스텔_전국 매매가격지수.pdf",
    "detached": "단독_전국 매매가격지수.pdf",
}
# 상가통합(502)에는 수익률 행이 없다. 공표된 세 유형을 평균하지 않는다.
YIELD_SHEETS = {
    "114": "office",
    "208": "mid_retail",
    "308": "small_retail",
    "408": "strata",
}
YIELD_ITEM = {
    "소득수익률(%)": "income",
    "자본수익률(%)": "capital",
    "투자수익률(%)": "investment",
}
HOUSING = ("apartment", "rowhouse", "officetel", "detached")
# KODEX KOSPI TR(359210) 연간 수익률(%). KRX KOSPI TR 원지수가 아니다.
KODEX_KOSPI_TR_PCT = {
    2021: 4.95,
    2022: -23.27,
    2023: 20.15,
    2024: -7.42,
    2025: 79.99,
}


def _r(v: float | None, nd: int = 4) -> float | None:
    if v is None:
        return None
    return round(float(v), nd)


def _series(values: dict[int, float | None]) -> list[dict[str, Any]]:
    return [{"year": y, "v": _r(values.get(y))} for y in YEARS]


def load_price_returns() -> dict[str, Any]:
    levels: dict[str, dict[int, float]] = {}
    returns: dict[str, dict[int, float | None]] = {}
    for key, name in PRICE_PDF.items():
        path = DATA / name
        idx = december_index(path)
        levels[key] = {y: idx[y] for y in sorted(idx) if 2020 <= y <= 2025}
        returns[key] = {y: december_return_pct(idx, y) for y in YEARS}
    return {"december": levels, "return_pct": returns}


def _latest_xlsx() -> Path:
    hits = list(SANGKWON_DIR.glob("*.xlsx"))
    if not hits:
        raise FileNotFoundError(f"xlsx 없음: {SANGKWON_DIR}")
    hits.sort(key=lambda p: p.stat().st_mtime)
    return hits[-1]


def load_commercial(path: Path | None = None) -> dict[str, Any]:
    xlsx = path or _latest_xlsx()
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    out: dict[str, dict[str, dict[int, float | None]]] = {}
    try:
        for sheet, kind in YIELD_SHEETS.items():
            ws = wb[sheet]
            quarters: list[tuple[int, int, int]] = []
            bucket: dict[str, dict[int, dict[int, float]]] = {
                m: {} for m in ("income", "capital", "investment")
            }
            for row in ws.iter_rows(values_only=True):
                cells = list(row)
                if not quarters:
                    quarters = [
                        (i, q[0], q[1])
                        for i, raw in enumerate(cells)
                        if (q := parse_quarter_header(raw))
                    ]
                    continue
                region = str(cells[0] or "").strip()
                name = str(cells[1] or "").strip() if len(cells) > 1 else ""
                item = str(cells[2] or "").strip() if len(cells) > 2 else ""
                metric = YIELD_ITEM.get(item)
                if region != "전국" or name != "합계" or metric is None:
                    continue
                for col, year, qtr in quarters:
                    if year not in YEARS and year != YEARS[0]:
                        if year < YEARS[0] or year > YEARS[-1]:
                            continue
                    val = parse_number(cells[col] if col < len(cells) else None)
                    if val is None:
                        continue
                    bucket[metric].setdefault(year, {})[qtr] = float(val)
            annual: dict[str, dict[int, float | None]] = {}
            for metric, by_year in bucket.items():
                annual[metric] = {}
                for year in YEARS:
                    annual[metric][year] = compound_annual(by_year.get(year, {}))
            out[kind] = annual
    finally:
        wb.close()
    return {"file": xlsx.name, "series": out}


def _year_values(block: dict[str, Any] | None) -> dict[int, float]:
    out: dict[int, float] = {}
    if not block:
        return out
    for p in block.get("points") or []:
        out[int(p["year"])] = float(p["v"])
    return out


def load_rates_and_kospi() -> dict[str, Any]:
    rate_path = _latest(ANNUAL_DIR, "시장금리*.csv")
    base_path = _latest(ANNUAL_DIR, "한국은행 기준금리*.csv")
    stock_path = _latest(ANNUAL_DIR, "주식시장*.csv")
    rates = parse_ecos_calendar_years(rate_path)
    base = parse_ecos_calendar_years(base_path)
    stock = parse_ecos_calendar_years(stock_path)
    cd = _year_values(find_named_series(rates, "CD(91일)"))
    ktb = _year_values(find_named_series(rates, "국고채(3년)"))
    base_rate = _year_values(find_named_series(base, "한국은행 기준금리"))
    close = _year_values(find_named_series(stock, "KOSPI_종가"))
    div_block = find_named_series(stock, "KOSPI_배당수익률")
    if div_block is None:
        for name, block in stock["series"].items():
            if name.replace(" ", "").startswith("KOSPI_배당수익률"):
                div_block = block
                break
    div = _year_values(div_block)
    # 2024 연 열 종가 2,399.49는 그해 12월 30일 종가와 같다. 평균 열이 아니다.
    if close.get(2024) is None or abs(close[2024] - 2399.49) > 0.02:
        raise RuntimeError(
            f"KOSPI_종가 연 열이 연말이 아님: 2024={close.get(2024)}"
        )
    price_ret = {
        y: ((close[y] / close[y - 1] - 1.0) * 100.0)
        if close.get(y) is not None and close.get(y - 1)
        else None
        for y in YEARS
    }
    return {
        "files": {
            "rates": rate_path.name,
            "base": base_path.name,
            "stock": stock_path.name,
        },
        "cd91": {y: cd.get(y) for y in YEARS},
        "ktb3": {y: ktb.get(y) for y in YEARS},
        "base_rate": {y: base_rate.get(y) for y in YEARS},
        "kospi_close": {y: close.get(y) for y in (2020,) + YEARS},
        "kospi_price_return_pct": price_ret,
        "kospi_dividend_yield_pct": {y: div.get(y) for y in YEARS},
    }


def imputed_monthly_per_m2(monthly: float, deposit: float, r_pct: float) -> float:
    """월세환산 ㎡당. r는 % (mean_simple). 전세는 넣지 않는다."""
    return float(monthly) + float(deposit) * (float(r_pct) / 100.0) / 12.0


def _median_rows(engine, sql: str, params: dict[str, Any]) -> dict[int, dict[str, float]]:
    with engine.connect() as conn:
        rows = conn.execute(text(sql), params).mappings()
        return {
            int(r["y"]): {"n": int(r["n"]), "med": float(r["med"])}
            for r in rows
            if r["med"] is not None
        }


def load_residential() -> dict[str, Any]:
    rent_eng = get_rent_engine()
    col_eng = get_collective_engine()
    built_eng = get_built_engine()
    if rent_eng is None or col_eng is None or built_eng is None:
        raise RuntimeError("RENT/COLLECTIVE/BUILT 데이터베이스 URL 없음")
    rent_sql = """
        SELECT contract_year AS y,
               COUNT(*) AS n,
               percentile_cont(0.5) WITHIN GROUP (ORDER BY monthly_per_m2) AS med
        FROM rent_transactions
        WHERE is_valid
          AND asset_type = :asset
          AND contract_year BETWEEN 2021 AND 2025
          AND monthly_rent_manwon > 0
          AND monthly_per_m2 > 0
        GROUP BY contract_year
    """
    sale_sql = """
        SELECT contract_year AS y,
               COUNT(*) AS n,
               percentile_cont(0.5) WITHIN GROUP (ORDER BY unit_price) AS med
        FROM collective_transactions
        WHERE is_valid
          AND asset_type = :asset
          AND contract_year BETWEEN 2021 AND 2025
          AND unit_price > 0
        GROUP BY contract_year
    """
    detached_sql = """
        SELECT contract_year AS y,
               COUNT(*) AS n,
               percentile_cont(0.5) WITHIN GROUP (ORDER BY price / gross_area) AS med
        FROM built_transactions
        WHERE is_valid
          AND asset_type = 'detached'
          AND contract_year BETWEEN 2021 AND 2025
          AND price > 0
          AND gross_area > 0
        GROUP BY contract_year
    """
    rent = {a: _median_rows(rent_eng, rent_sql, {"asset": a}) for a in HOUSING}
    sale = {a: _median_rows(col_eng, sale_sql, {"asset": a}) for a in HOUSING if a != "detached"}
    sale["detached"] = _median_rows(built_eng, detached_sql, {})
    r_sql = """
        WITH b AS (
            SELECT contract_year AS y,
                   percentile_cont(0.5) WITHIN GROUP (ORDER BY deposit_per_m2)
                       FILTER (WHERE COALESCE(monthly_rent_manwon, 0) = 0 AND deposit_per_m2 > 0) AS j,
                   COUNT(*) FILTER (
                       WHERE COALESCE(monthly_rent_manwon, 0) = 0 AND deposit_per_m2 > 0
                   ) AS n_j,
                   percentile_cont(0.5) WITHIN GROUP (ORDER BY deposit_per_m2)
                       FILTER (WHERE monthly_rent_manwon > 0 AND deposit_per_m2 > 0) AS d,
                   percentile_cont(0.5) WITHIN GROUP (ORDER BY monthly_per_m2)
                       FILTER (
                           WHERE monthly_rent_manwon > 0 AND deposit_per_m2 > 0 AND monthly_per_m2 > 0
                       ) AS m,
                   COUNT(*) FILTER (
                       WHERE monthly_rent_manwon > 0 AND deposit_per_m2 > 0 AND monthly_per_m2 > 0
                   ) AS n_m
            FROM rent_transactions
            WHERE is_valid
              AND asset_type = :asset
              AND contract_year BETWEEN 2021 AND 2025
              AND building_key IS NOT NULL
              AND btrim(building_key) <> ''
            GROUP BY contract_year, building_key
        )
        SELECT y,
               COUNT(*) FILTER (
                   WHERE n_j >= 3 AND n_m >= 3 AND j > d AND m > 0
                     AND (12.0 * m / (j - d) * 100.0) BETWEEN 1 AND 15
               ) AS n,
               AVG(12.0 * m / (j - d) * 100.0) FILTER (
                   WHERE n_j >= 3 AND n_m >= 3 AND j > d AND m > 0
                     AND (12.0 * m / (j - d) * 100.0) BETWEEN 1 AND 15
               ) AS med
        FROM b
        GROUP BY y
    """
    conv_sql = """
        SELECT contract_year AS y,
               COUNT(*) AS n,
               percentile_cont(0.5) WITHIN GROUP (
                   ORDER BY monthly_per_m2 + COALESCE(deposit_per_m2, 0) * (:r / 100.0) / 12.0
               ) AS med
        FROM rent_transactions
        WHERE is_valid
          AND asset_type = :asset
          AND contract_year = :year
          AND monthly_rent_manwon > 0
          AND monthly_per_m2 > 0
        GROUP BY contract_year
    """
    rates = {a: _median_rows(rent_eng, r_sql, {"asset": a}) for a in HOUSING}
    converted: dict[str, dict[int, dict[str, float]]] = {a: {} for a in HOUSING}
    for asset in HOUSING:
        for year in YEARS:
            rate = rates[asset].get(year)
            if rate is None or rate["n"] < 1:
                continue
            row = _median_rows(
                rent_eng,
                conv_sql,
                {"asset": asset, "year": year, "r": rate["med"]},
            ).get(year)
            if row is not None:
                converted[asset][year] = row
    income: dict[str, dict[int, float | None]] = {}
    converted_income: dict[str, dict[int, float | None]] = {}
    counts: dict[str, dict[int, dict[str, int | None]]] = {}
    for asset in HOUSING:
        income[asset] = {}
        converted_income[asset] = {}
        counts[asset] = {}
        for year in YEARS:
            r = rent[asset].get(year)
            s = sale[asset].get(year)
            rate = rates[asset].get(year)
            conv = converted[asset].get(year)
            n_rent = int(r["n"]) if r else None
            n_sale = int(s["n"]) if s else None
            counts[asset][year] = {
                "n_rent": n_rent,
                "n_sale": n_sale,
                "n_buildings": int(rate["n"]) if rate else None,
                "r_pct": float(rate["med"]) if rate else None,
            }
            sale_ok = s is not None and s["med"] > 0 and (n_sale or 0) >= MIN_N
            if r is None or not sale_ok or r["med"] <= 0 or (n_rent or 0) < MIN_N:
                income[asset][year] = None
            else:
                income[asset][year] = 100.0 * (r["med"] * 12.0) / s["med"]
            if conv is None or not sale_ok or conv["med"] <= 0 or rate is None or (conv["n"] or 0) < MIN_N:
                converted_income[asset][year] = None
            else:
                converted_income[asset][year] = 100.0 * (conv["med"] * 12.0) / s["med"]
    return {
        "income_pct": income,
        "converted_income_pct": converted_income,
        "counts": counts,
        "min_n": MIN_N,
    }


def _add_pct(
    left: dict[int, float | None],
    right: dict[int, float | None],
) -> dict[int, float | None]:
    """주거 연간 투자 = 그해 소득 + 그해 자본. 분기로 나누지 않는다."""
    return {
        y: (left[y] + right[y]) if left.get(y) is not None and right.get(y) is not None else None
        for y in YEARS
    }


def _spread(income: dict[int, float | None], ktb: dict[int, float | None]) -> dict[int, float | None]:
    return {
        y: (income[y] - ktb[y]) if income.get(y) is not None and ktb.get(y) is not None else None
        for y in YEARS
    }


def build_snapshot() -> dict[str, Any]:
    price = load_price_returns()
    commercial = load_commercial()
    market = load_rates_and_kospi()
    housing = load_residential()
    apt = housing["income_pct"]["apartment"]
    apt_conv = housing["converted_income_pct"]["apartment"]
    apt10 = {
        y: (apt[y] * (1.0 - EXPENSE_RATIO)) if apt.get(y) is not None else None
        for y in YEARS
    }
    apt10_conv = {
        y: (apt_conv[y] * (1.0 - EXPENSE_RATIO)) if apt_conv.get(y) is not None else None
        for y in YEARS
    }
    comm = commercial["series"]
    ktb = market["ktb3"]
    income_rows = {
        "office": comm["office"]["income"],
        "mid_retail": comm["mid_retail"]["income"],
        "small_retail": comm["small_retail"]["income"],
        "strata": comm["strata"]["income"],
        "apartment": apt,
        "apartment_converted": apt_conv,
        "rowhouse": housing["income_pct"]["rowhouse"],
        "rowhouse_converted": housing["converted_income_pct"]["rowhouse"],
        "officetel": housing["income_pct"]["officetel"],
        "officetel_converted": housing["converted_income_pct"]["officetel"],
        "detached": housing["income_pct"]["detached"],
        "detached_converted": housing["converted_income_pct"]["detached"],
        "apartment_expense10": apt10,
        "apartment_converted_expense10": apt10_conv,
        "ktb3": ktb,
        "cd91": market["cd91"],
        "base_rate": market["base_rate"],
    }
    spreads = {
        "office": _spread(comm["office"]["income"], ktb),
        "mid_retail": _spread(comm["mid_retail"]["income"], ktb),
        "small_retail": _spread(comm["small_retail"]["income"], ktb),
        "strata": _spread(comm["strata"]["income"], ktb),
        "apartment": _spread(apt, ktb),
        "apartment_converted": _spread(apt_conv, ktb),
        "apartment_expense10": _spread(apt10, ktb),
        "apartment_converted_expense10": _spread(apt10_conv, ktb),
    }
    capital_rows = {
        "office": comm["office"]["capital"],
        "mid_retail": comm["mid_retail"]["capital"],
        "small_retail": comm["small_retail"]["capital"],
        "strata": comm["strata"]["capital"],
        "apartment": price["return_pct"]["apartment"],
        "rowhouse": price["return_pct"]["rowhouse"],
        "officetel": price["return_pct"]["officetel"],
        "detached": price["return_pct"]["detached"],
        "kospi_price": market["kospi_price_return_pct"],
    }
    investment_rows = {
        "office": comm["office"]["investment"],
        "mid_retail": comm["mid_retail"]["investment"],
        "small_retail": comm["small_retail"]["investment"],
        "strata": comm["strata"]["investment"],
        "apartment": _add_pct(apt, capital_rows["apartment"]),
        "apartment_converted": _add_pct(apt_conv, capital_rows["apartment"]),
        "apartment_expense10": _add_pct(apt10, capital_rows["apartment"]),
        "apartment_converted_expense10": _add_pct(apt10_conv, capital_rows["apartment"]),
        "rowhouse": _add_pct(housing["income_pct"]["rowhouse"], capital_rows["rowhouse"]),
        "rowhouse_converted": _add_pct(housing["converted_income_pct"]["rowhouse"], capital_rows["rowhouse"]),
        "officetel": _add_pct(housing["income_pct"]["officetel"], capital_rows["officetel"]),
        "officetel_converted": _add_pct(housing["converted_income_pct"]["officetel"], capital_rows["officetel"]),
        "detached": _add_pct(housing["income_pct"]["detached"], capital_rows["detached"]),
        "detached_converted": _add_pct(housing["converted_income_pct"]["detached"], capital_rows["detached"]),
        "kospi_tr": {y: KODEX_KOSPI_TR_PCT[y] for y in YEARS},
    }
    return {
        "question": "최근 5년, 상업 수익률과 아파트 임대료/매매가는 금리·주식과 어디서 나란한가",
        "years": list(YEARS),
        "scope": "전국",
        "residential_investment": "annual_income_plus_capital",
        "notes": [
            "상업 연간은 그해 1–4분기 복리. 전국·합계 행만. 상권 평균 없음.",
            "상가통합 시트에는 수익률 행이 없다. 중대형·소규모·집합을 평균하지 않고 각각 둔다.",
            "주거 현금 소득은 월세>0 계약의 ㎡당 월세 가운데값×12 / 매매 ㎡당 가운데값. 전세는 제외.",
            "주거 월세환산은 같은 계약에 보증금×그해 단순평균 전환율/12를 더한다. 전환율은 국고채가 아니다. 전세 계약은 넣지 않는다.",
            "아파트 10% 선은 소득의 0.9배 가정. 공표 순영업소득이 아님.",
            "주거 자본은 매매가격지수 12월/전년 12월−1. 주거 투자는 그해 소득+그해 자본. 분기로 나누지 않음.",
            "단독 임대 면적은 계약면적, 단독 매매 면적은 연면적.",
            "KOSPI_종가 연 열은 연말 종가. 2024=2399.49. 자본 칸의 가격수익률만 만든다. 배당수익률은 소득 칸의 수준. 둘을 더하지 않음.",
            "투자 칸의 코스피는 KODEX KOSPI TR(359210) 연간 수익률. KRX KOSPI TR 원지수가 아님. 보수·추적오차가 있다.",
            f"주거 소득은 임대 건수 또는 매매 건수가 {MIN_N} 미만이면 빈칸.",
        ],
        "sources": {
            "commercial_xlsx": commercial["file"],
            "rates": market["files"]["rates"],
            "base_rate": market["files"]["base"],
            "stock": market["files"]["stock"],
            "price_pdf": PRICE_PDF,
        },
        "income_pct": {k: _series(v) for k, v in income_rows.items()},
        "spread_vs_ktb3_pct": {k: _series(v) for k, v in spreads.items()},
        "capital_pct": {k: _series(v) for k, v in capital_rows.items()},
        "investment_pct": {k: _series(v) for k, v in investment_rows.items()},
        "kospi_dividend_yield_pct": _series(market["kospi_dividend_yield_pct"]),
        "price_december": {
            k: [{"year": y, "v": _r(levels.get(y), 2)} for y in range(2020, 2026) if y in levels]
            for k, levels in price["december"].items()
        },
        "residential_counts": {
            asset: [
                {
                    "year": y,
                    "n_rent": counts[y]["n_rent"],
                    "n_sale": counts[y]["n_sale"],
                    "n_buildings": counts[y].get("n_buildings"),
                    "r_pct": _r(counts[y].get("r_pct"), 2),
                }
                for y in YEARS
            ]
            for asset, counts in housing["counts"].items()
        },
    }


def main() -> None:
    snap = build_snapshot()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(snap, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUT)


if __name__ == "__main__":
    main()
