#!/usr/bin/env python3
"""토스페이먼츠 결제경로 PPT 생성 — ch2data.com 화면 캡처 포함."""

from __future__ import annotations

import asyncio
import json
import os
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from playwright.async_api import async_playwright

REPO = Path(__file__).resolve().parents[1]
OUT_DIR = REPO / "hub" / "toss-review"
SHOT_DIR = OUT_DIR / "screenshots"
PPT_PATH = OUT_DIR / "CH2DATA_결제경로.pptx"

MERCHANT = {
    "name": "씨에이치투",
    "biz_no": "607-23-96932",
    "url": "https://ch2data.com/",
    "test_id": os.environ.get("PLATFORM_REVIEW_EMAIL", "review@ch2data.com"),
    "test_pw": os.environ.get("PLATFORM_REVIEW_PASSWORD", ""),
    "ceo": "정탁영",
    "address": "충청북도 청주시 흥덕구 서현중로35번길 25, 401호(가경동)",
    "phone": "010-5801-7953",
    "email": "ch2data.dev@gmail.com",
}

BLUE = RGBColor(0x31, 0x82, 0xCE)
BLACK = RGBColor(0x11, 0x18, 0x27)
GRAY = RGBColor(0x64, 0x74, 0x8B)


def _add_title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(8.4), Inches(4.5))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "① 가맹점 정보"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = BLUE

    lines = [
        f"(1) 상호명: {MERCHANT['name']}",
        f"(2) 사업자번호: {MERCHANT['biz_no']}",
        f"(3) URL: {MERCHANT['url']}",
        f"(4) Test ID: {MERCHANT['test_id']}",
        f"(5) Test PW: {MERCHANT['test_pw']}",
    ]
    for line in lines:
        para = tf.add_paragraph()
        para.text = line
        para.font.size = Pt(20)
        para.font.color.rgb = BLACK
        para.space_before = Pt(10)


def _add_section_slide(prs: Presentation, title: str, subtitle: str, image_path: Path, url: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.6))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(22)
    tp.font.bold = True
    tp.font.color.rgb = BLUE

    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.8), Inches(9), Inches(0.4))
    sp = sub_box.text_frame.paragraphs[0]
    sp.text = subtitle
    sp.font.size = Pt(12)
    sp.font.color.rgb = GRAY

    if image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(0.5), Inches(1.2), width=Inches(9))

    cap = slide.shapes.add_textbox(Inches(0.5), Inches(6.85), Inches(9), Inches(0.35))
    cp = cap.text_frame.paragraphs[0]
    cp.text = f"URL: {url}"
    cp.font.size = Pt(10)
    cp.font.color.rgb = GRAY


def _add_notes_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.6), Inches(8.6), Inches(6))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "결제경로 제작 유의사항"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = BLUE
    notes = [
        "PPT 형식으로 제작 (본 파일)",
        "캡처 화면에 URL을 슬라이드 하단에 표기",
        "무형상품(월 구독 SaaS) — CH2 Macro / Macro+FieldNote",
        "카드 자동결제(빌링) — 등록창 + 매월 자동 청구",
        "로그인 후 구매 — 심사용 이메일·비밀번호 계정 제공",
        "테스트 키로 카드 등록·결제창 연동 (본 PPT ⑥)",
    ]
    for note in notes:
        para = tf.add_paragraph()
        para.text = f"• {note}"
        para.font.size = Pt(16)
        para.font.color.rgb = BLACK
        para.space_before = Pt(8)


async def capture_screenshots() -> dict[str, Path]:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    SHOT_DIR.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}

    async with async_playwright() as p:
        browser = await p.chromium.launch()
        context = await browser.new_context(viewport={"width": 1360, "height": 900})
        page = await context.new_page()

        async def shot(name: str, url: str, *, full_page: bool = True, before=None):
            await page.goto(url, wait_until="networkidle")
            if before:
                await before()
            target = SHOT_DIR / f"{name}.png"
            await page.screenshot(path=str(target), full_page=full_page)
            paths[name] = target

        await shot("01_home_footer", "https://ch2data.com/")
        await page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
        await page.screenshot(path=str(SHOT_DIR / "02_footer.png"), full_page=False)
        paths["02_footer"] = SHOT_DIR / "02_footer.png"

        await shot("03_refund", "https://ch2data.com/refund/")
        await shot("04_subscription", "https://ch2data.com/subscription/")

        async def login_review():
            await context.request.post(
                "https://ch2data.com/api/auth/review-login",
                data=json.dumps(
                    {
                        "email": MERCHANT["test_id"],
                        "password": MERCHANT["test_pw"],
                    }
                ),
                headers={"Content-Type": "application/json"},
            )
            await page.goto("https://ch2data.com/subscribe/", wait_until="networkidle")
            await page.wait_for_timeout(1200)

        async def show_login_form():
            await page.wait_for_selector("#review-login", state="attached", timeout=15000)
            await page.evaluate(
                """() => {
                  const box = document.getElementById('review-login');
                  if (box) box.hidden = false;
                }"""
            )

        await shot("05_login", "https://ch2data.com/subscribe/", before=show_login_form)
        await shot("06_subscribe_logged_in", "https://ch2data.com/subscribe/", before=login_review)

        await page.goto("https://ch2data.com/subscribe/", wait_until="networkidle")
        await login_review()
        card_btn = page.locator(".btn-pay--card").first
        if await card_btn.count() > 0 and not await card_btn.is_disabled():
            await card_btn.click()
            await page.wait_for_timeout(3500)
            target = SHOT_DIR / "07_payment_window.png"
            await page.screenshot(path=str(target), full_page=False)
            paths["07_payment_window"] = target
        else:
            note = SHOT_DIR / "07_payment_note.png"
            await page.screenshot(path=str(note), full_page=False)
            paths["07_payment_window"] = note

        await browser.close()
    return paths


def build_ppt(paths: dict[str, Path]) -> None:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    _add_title_slide(prs)
    _add_notes_slide(prs)
    _add_section_slide(
        prs,
        "② 하단 정보",
        f"홈페이지 하단 사업자 정보 (캡처 시각: {stamp})",
        paths.get("02_footer", paths.get("01_home_footer", Path())),
        "https://ch2data.com/",
    )
    _add_section_slide(
        prs,
        "③ 환불규정 (무형상품)",
        "디지털 콘텐츠 월 구독 환불 정책",
        paths.get("03_refund", Path()),
        "https://ch2data.com/refund/",
    )
    _add_section_slide(
        prs,
        "④ 로그인",
        "심사용 이메일·비밀번호 로그인 (구독 페이지)",
        paths.get("05_login", Path()),
        "https://ch2data.com/subscribe/",
    )
    _add_section_slide(
        prs,
        "⑤ 상품·요금",
        "구독 상품명·금액·설명",
        paths.get("04_subscription", Path()),
        "https://ch2data.com/subscription/",
    )
    _add_section_slide(
        prs,
        "⑤ 구매 과정",
        "로그인 후 월 구독 · 카드 자동결제 시작",
        paths.get("06_subscribe_logged_in", Path()),
        "https://ch2data.com/subscribe/",
    )
    _add_section_slide(
        prs,
        "⑥ 카드 결제경로",
        "「카드로 구독하기」 클릭 후 토스 카드 등록·결제창 (테스트 키)",
        paths.get("07_payment_window", Path()),
        "https://ch2data.com/subscribe/",
    )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(str(PPT_PATH))


async def main() -> None:
    paths = await capture_screenshots()
    build_ppt(paths)
    print(f"OK: {PPT_PATH}")


if __name__ == "__main__":
    asyncio.run(main())
