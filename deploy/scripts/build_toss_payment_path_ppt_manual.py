#!/usr/bin/env python3
"""deploy/hub/toss-review/screenshots PNG로 토스 결제경로 PPT 생성."""

from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
ROOT = REPO.parent
OUT_DIR = REPO / "hub" / "toss-review"
SHOT_DIR = OUT_DIR / "screenshots"
PPT_PATH = OUT_DIR / "CH2DATA_payment_path.pptx"
PPT_PATH_KO = OUT_DIR / "CH2DATA_결제경로.pptx"
LOCAL_ENV = ROOT / "deploy" / "local" / "toss-test-keys.env"

BLUE = RGBColor(0x31, 0x82, 0xCE)
BLACK = RGBColor(0x11, 0x18, 0x27)
GRAY = RGBColor(0x64, 0x74, 0x8B)

SLIDE_W = Inches(10)
SLIDE_H = Inches(7.5)
IMG_LEFT = Inches(0.35)
IMG_TOP = Inches(1.05)
IMG_MAX_W = Inches(9.3)
IMG_MAX_H = Inches(6.35)

SLIDES = [
    ("② 하단 정보", "홈페이지 하단 사업자 정보", "02_footer.png"),
    ("③ 환불규정 (무형상품)", "디지털 콘텐츠 월 구독 환불 정책", "03_refund.png"),
    ("④ 로그인", "심사용 이메일·비밀번호 로그인", "04_login.png"),
    ("⑤ 구독 정책 (1/2)", "상품·요금·무료 체험", "05_subscription_1.png"),
    ("⑤ 구독 정책 (2/2)", "제품 추가·해제·결제 채널·해지", "05_subscription_2.png"),
    ("⑤ 구매 과정", "로그인 후 구독 페이지 (카드 결제창·심사용)", "06_subscribe.png"),
    (
        "⑥ 카드 결제경로",
        "테스트 키 · 카드 결제창(심사용) → 토스 결제창 (실제 결제 없음)",
        "07_payment_window.png",
    ),
]


def _load_env_file(path: Path) -> None:
    if not path.is_file():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, val = line.split("=", 1)
        key, val = key.strip(), val.strip()
        if key and val:
            os.environ.setdefault(key, val)


def merchant_info() -> dict[str, str]:
    _load_env_file(LOCAL_ENV)
    return {
        "name": "씨에이치투",
        "biz_no": "607-23-96932",
        "url": "https://ch2data.com/",
        "test_id": os.environ.get("PLATFORM_REVIEW_EMAIL", "review@ch2data.com"),
        "test_pw": os.environ.get("PLATFORM_REVIEW_PASSWORD", "(VPS PLATFORM_REVIEW_PASSWORD)"),
    }


def _fit_picture(slide, image_path: Path):
    pic = slide.shapes.add_picture(str(image_path), IMG_LEFT, IMG_TOP)
    ratio = pic.width / pic.height
    if pic.width > IMG_MAX_W:
        pic.width = int(IMG_MAX_W)
        pic.height = int(IMG_MAX_W / ratio)
    if pic.height > IMG_MAX_H:
        pic.height = int(IMG_MAX_H)
        pic.width = int(IMG_MAX_H * ratio)
    pic.left = int((SLIDE_W - pic.width) / 2)
    pic.top = int(IMG_TOP)


def _add_title_slide(prs: Presentation, merchant: dict[str, str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(8.4), Inches(4.5))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "① 가맹점 정보"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = BLUE
    for line in [
        f"(1) 상호명: {merchant['name']}",
        f"(2) 사업자번호: {merchant['biz_no']}",
        f"(3) URL: {merchant['url']}",
        f"(4) Test ID: {merchant['test_id']}",
        f"(5) Test PW: {merchant['test_pw']}",
        "(6) MID: ch2datoe5e (API 개별 연동 · 테스트 키)",
    ]:
        para = tf.add_paragraph()
        para.text = line
        para.font.size = Pt(18 if line.startswith("(6)") else 20)
        para.font.color.rgb = BLACK
        para.space_before = Pt(8)


def _add_notes_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.8), Inches(8.6), Inches(6))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "결제경로 유의사항"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = BLUE
    for line in [
        "월 구독 SaaS (Macro / Macro+FieldNote) — 무형상품",
        "테스트 키 연동 · 카드 결제창(심사용) 캡처 포함",
        "운영: 카드 자동결제(빌링) — 가맹 빌링 계약 후 활성화",
        "각 슬라이드 하단 URL: https://ch2data.com/subscribe/",
    ]:
        para = tf.add_paragraph()
        para.text = f"• {line}"
        para.font.size = Pt(16)
        para.font.color.rgb = BLACK
        para.space_before = Pt(8)


def _add_image_slide(prs: Presentation, title: str, subtitle: str, image_name: str) -> None:
    image_path = SHOT_DIR / image_name
    if not image_path.is_file():
        raise FileNotFoundError(f"캡처 없음: {image_path}")
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9.2), Inches(0.45))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(20)
    tp.font.bold = True
    tp.font.color.rgb = BLUE
    sub_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.62), Inches(9.2), Inches(0.35))
    sp = sub_box.text_frame.paragraphs[0]
    sp.text = subtitle
    sp.font.size = Pt(11)
    sp.font.color.rgb = GRAY
    _fit_picture(slide, image_path)
    cap = slide.shapes.add_textbox(Inches(0.4), Inches(7.05), Inches(9.2), Inches(0.3))
    cp = cap.text_frame.paragraphs[0]
    url = "https://ch2data.com/subscribe/" if "⑥" in title or "⑤ 구매" in title else "https://ch2data.com/"
    cp.text = f"URL: {url}"
    cp.font.size = Pt(9)
    cp.font.color.rgb = GRAY


def build() -> None:
    merchant = merchant_info()
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    _add_title_slide(prs, merchant)
    _add_notes_slide(prs)
    for title, subtitle, fname in SLIDES:
        _add_image_slide(prs, title, subtitle, fname)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(str(PPT_PATH))
    prs.save(str(PPT_PATH_KO))
    print(f"OK: {PPT_PATH}")
    print(f"OK: {PPT_PATH_KO} ({datetime.now():%Y-%m-%d %H:%M})")


if __name__ == "__main__":
    build()
